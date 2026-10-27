"""Real source files retain the cited passage's filename and exact location."""
from __future__ import annotations

import numpy as np
import pymupdf
import pytest

from backend.db.document_store import DocumentStore
from backend.ingestion.markdown_processor import chunk_markdown
from backend.ingestion.office_processor import process_office
from backend.ingestion.pdf_processor import process_pdf
from backend.retrieval.vector_retriever import VectorIndex, retrieve_dense
from backend.schemas import Chunk, CitationSpan, GenerationResult, MAX_CHUNK_TEXT_CHARACTERS, ReflectionVerdict, RetrievalResult
from backend.utils.provenance import evidence_metadata
from backend.utils.uploads import remap_chunks, source_identity
from backend.tests.test_ingest_safety import api_state


class PassageEmbedder:
    """Deterministic retrieval isolates citation mapping from model quality."""
    dimension = 2
    model_name = "citation-fixture"

    def __init__(self, needle): self.needle = needle
    def embed(self, texts):
        return np.array([[1, 0] if self.needle in text else [0, 1] for text in texts], dtype=np.float32)
    def embed_one(self, text): return self.embed([text])[0]
    embed_query = embed_one


def _query_passage(chunks, filename, needle, tmp_path, monkeypatch):
    monkeypatch.setenv("LUMINRAG_SKIP_DOTENV", "1")
    from backend import main
    chunks = remap_chunks(chunks, source_identity(filename), filename, f"source/{filename}")
    index = VectorIndex(PassageEmbedder(needle))
    index.build(chunks)
    with DocumentStore(tmp_path / "citations.db") as store:
        store.save_chunks(chunks)
        retrieved = retrieve_dense(needle, index, store, {"vector_retriever": {"top_k": 1}})
        assert needle in retrieved.chunks[0].text
        return main._chunks_to_evidence(retrieved)[0]


def test_evidence_returns_complete_passage_after_500_characters(monkeypatch):
    monkeypatch.setenv("LUMINRAG_SKIP_DOTENV", "1")
    from backend import main
    text = "Introductory material. " * 50 + "The exact cited conclusion is here."
    chunk = Chunk(id="passage", source_id="stable-source", modality="pdf", text=text, metadata={"original_filename": "lecture.pdf", "page_start": 4, "page_end": 5, "section_title": "Energy transfer"})
    evidence = main._chunks_to_evidence(RetrievalResult(chunks=[chunk], subgraph=[], routing_mode="dense"))[0]
    assert evidence.text == text
    assert "exact cited conclusion" in evidence.text
    assert evidence.source == "stable-source"
    assert evidence.source_filename == "lecture.pdf"
    assert (evidence.page, evidence.page_end, evidence.section_title) == (4, 5, "Energy transfer")


def test_legacy_huge_passage_is_bounded_by_shared_chunk_limit():
    chunk = Chunk(id="large", source_id="source", modality="pdf", text="x" * (MAX_CHUNK_TEXT_CHARACTERS + 1))
    assert len(evidence_metadata(chunk)["text"]) == MAX_CHUNK_TEXT_CHARACTERS


def test_real_pdf_later_page_query_uses_actual_passage_page(tmp_path, monkeypatch):
    needle = "LATE_PAGE_CITATION"
    path = tmp_path / "citation-course.pdf"
    with pymupdf.open() as document:
        for page_number in range(1, 4):
            page = document.new_page()
            if page_number == 1:
                page.insert_text((50, 50), "Course notes", fontsize=18)
            text = "\n".join(f"Page {page_number} educational source sentence {line_number:02d}." for line_number in range(28))
            if page_number == 3:
                text += f"\n{needle}: this mechanism appears on the third page."
            assert page.insert_textbox(pymupdf.Rect(50, 90, 545, 750), text, fontsize=10) >= 0
        document.save(path)
    chunks = process_pdf(path, {"pdf_processor": {"max_chunk_chars": 380, "min_heading_size": 14}, "data": {"processed_dir": str(tmp_path / "processed")}})
    assert len(chunks) > 3
    evidence = _query_passage(chunks, path.name, needle, tmp_path, monkeypatch)
    assert evidence.page == 3
    assert evidence.page_end == 3
    assert evidence.source_filename == path.name
    assert evidence.section_title == "Course notes"
    assert needle in evidence.text


def test_selected_sentence_narrows_a_split_pdf_chunk_to_its_actual_page(tmp_path, monkeypatch):
    monkeypatch.setenv("LUMINRAG_SKIP_DOTENV", "1")
    from backend import main
    quote = "The selected theorem belongs to the second page."
    path = tmp_path / "two-page-source.pdf"
    with pymupdf.open() as document:
        first = document.new_page()
        first.insert_text((50, 50), "Source chapter", fontsize=18)
        first.insert_text((50, 90), "Opening context establishes the background.", fontsize=10)
        second = document.new_page()
        text = quote + " Supporting detail" * 16
        assert second.insert_textbox(pymupdf.Rect(50, 90, 545, 750), text, fontsize=10) >= 0
        document.save(path)
    chunks = process_pdf(path, {"pdf_processor": {"max_chunk_chars": 200}, "data": {"processed_dir": str(tmp_path / "processed")}})
    chunk = next(chunk for chunk in chunks if quote in chunk.text)
    assert (chunk.metadata["page_start"], chunk.metadata["page_end"]) == (1, 2)
    assert chunk.metadata["page_spans"]
    start = chunk.text.index(quote)
    citation = CitationSpan(chunk_id=chunk.id, source_number=1, quote=quote, quote_start=start, quote_end=start + len(quote))
    with DocumentStore(tmp_path / "quoted.db") as store:
        store.save_chunks([chunk])
        restored = store.get_chunk(chunk.id)
        evidence = main._chunks_to_evidence(RetrievalResult(chunks=[restored], subgraph=[], routing_mode="dense"), [citation])[0]
    assert (evidence.page, evidence.page_end) == (2, 2)
    assert evidence.text == chunk.text
    assert evidence.quote == quote
    assert evidence.text[evidence.quote_start:evidence.quote_end] == quote


def test_query_response_wires_only_verified_generation_quote(api_state):
    from backend.utils.state import current_revision
    main = api_state.main
    chunk = api_state.old.model_copy(update={"text": "Previous context. The supporting sentence.", "metadata": {"page_start": 1, "page_end": 2, "page_spans": [[0, 18, 1], [18, 42, 2]], "original_filename": "course.pdf"}})
    quote = "The supporting sentence."
    start = chunk.text.index(quote)
    generated = GenerationResult(answer="An answer [1].", routing_mode="dense", evidence_chunk_ids=[chunk.id], reflection=ReflectionVerdict(needs_retrieval=True, is_relevant=True, is_supported=True, is_useful=True, reasoning="supported"), citations=[CitationSpan(chunk_id=chunk.id, source_number=1, quote=quote, quote_start=start, quote_end=start + len(quote))])
    retrieved = RetrievalResult(chunks=[chunk], subgraph=[], routing_mode="dense")
    prepared = main._PreparedQuery("question", retrieved, {}, current_revision())
    response = main._finish_query(main.QueryRequest(query="question"), prepared, generated)
    assert response.evidence[0].quote == quote
    assert response.evidence[0].page == 2
    assert response.evidence[0].source_filename == "course.pdf"


@pytest.mark.parametrize("quote,start,end", [("invented quote", 0, 5), ("source", 100, 106), ("source", 0, 6)])
def test_mismatched_quote_never_creates_a_false_highlight(quote, start, end):
    chunk = Chunk(id="source", source_id="source", modality="pdf", text="Actual source sentence.", metadata={"page_start": 2, "page_end": 3})
    citation = CitationSpan(chunk_id=chunk.id, source_number=1, quote=quote, quote_start=start, quote_end=end)
    payload = evidence_metadata(chunk, citation)
    assert payload["quote"] is None
    assert payload["quote_start"] is None and payload["quote_end"] is None
    assert (payload["page"], payload["page_end"]) == (2, 3)


def test_quote_offsets_are_unicode_code_points_and_legacy_pages_remain_truthful():
    text = "🌿 Plant energy comes from light."
    quote = "Plant energy"
    start = text.index(quote)
    chunk = Chunk(id="unicode", source_id="source", modality="pdf", text=text, metadata={"page_start": 4, "page_end": 6})
    payload = evidence_metadata(chunk, CitationSpan(chunk_id=chunk.id, source_number=1, quote=quote, quote_start=start, quote_end=start + len(quote)))
    assert payload["quote"] == quote and payload["quote_start"] == 2
    assert (payload["page"], payload["page_end"]) == (4, 6)


def test_real_pptx_later_slide_query_preserves_slide_number_and_title(tmp_path, monkeypatch):
    from pptx import Presentation
    needle = "LATE_SLIDE_CITATION"
    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Opening slide"
    first.placeholders[1].text = "A short introduction."
    presentation.slides.add_slide(presentation.slide_layouts[6])  # blank slide 2
    third = presentation.slides.add_slide(presentation.slide_layouts[1])
    third.shapes.title.text = "Detailed kinetics"
    third.placeholders[1].text = f"{needle}: the mechanism is explained on slide three."
    path = tmp_path / "lecture.pptx"
    presentation.save(path)
    chunks = process_office(path, {})
    assert len(chunks) == 2  # blank slide comments are never evidence
    assert all("<!-- Slide number:" not in chunk.text for chunk in chunks)
    evidence = _query_passage(chunks, path.name, needle, tmp_path, monkeypatch)
    assert evidence.slide_number == 3
    assert evidence.page is None
    assert evidence.source_filename == path.name
    assert evidence.title == "Detailed kinetics"
    assert evidence.section_title == "Detailed kinetics"
    assert needle in evidence.text


def test_markdown_comment_only_material_is_not_a_source_passage():
    assert chunk_markdown("<!-- commentary only -->\n\n<!-- another\ncomment -->", "source", {}) == []
    code = "```html\n<!-- Slide number: 12 -->\n```"
    chunks = chunk_markdown(code, "source", {})
    assert len(chunks) == 1
    assert "<!-- Slide number: 12 -->" in chunks[0].text
    assert "slide_number" not in chunks[0].metadata


@pytest.mark.parametrize("metadata", [
    {"page_start": float("inf"), "page_end": {}, "slide_number": [], "title": {}},
    {"page_start": -1, "page_end": 0, "slide_number": True, "source_filename": []},
    {"page_start": 2.5, "slide_number": "bad", "relevancy_score": float("nan"), "retrieval_source": {}},
])
def test_malformed_legacy_location_metadata_is_safely_omitted(metadata):
    payload = evidence_metadata(Chunk(id="source", source_id="source", modality="pdf", text="passage", metadata=metadata))
    assert payload["page"] is None
    assert payload["slide_number"] is None
    assert payload["relevancy_score"] is None
