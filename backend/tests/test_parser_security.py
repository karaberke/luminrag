"""Small controlled fixtures for parser resource boundaries; no model calls."""
from __future__ import annotations

import io
import json
from pathlib import Path
import sys
import time
from types import SimpleNamespace
from unittest.mock import MagicMock
import wave
import zipfile

import numpy as np
import pytest

from backend.ingestion import media_limits, office_processor
from backend.ingestion.audio_processor import _build_chunks
from backend.ingestion import video_transcriber
from backend.utils.processes import MIB, ProcessLimits, ProcessingLimitError, run_bounded


def workbook_path(tmp_path, *, rows=1000, columns=128, sheets=1):
    from openpyxl import Workbook
    workbook = Workbook()
    for index in range(sheets):
        sheet = workbook.active if index == 0 else workbook.create_sheet()
        sheet["A1"] = "Header"
        sheet.cell(row=rows, column=columns, value="Tail")
    path = tmp_path / "table.xlsx"
    workbook.save(path)
    return path


def small_wav(path):
    with wave.open(str(path), "wb") as output:
        output.setparams((1, 2, 16_000, 0, "NONE", "not compressed"))
        output.writeframes(b"\x00\x00" * 160)
    return path


def test_worker_diagnostics_cannot_corrupt_json_response(tmp_path):
    command = """
import os
from backend.ingestion import office_processor, parser_worker
def convert(path):
    print('Python parser diagnostic', flush=True)
    os.write(1, b'Native parser diagnostic\\n')
    return 'Converted document'
office_processor._convert_office = convert
parser_worker.main()
"""
    payload = json.dumps({"kind": "office", "path": str(tmp_path / "fixture.docx")}).encode()
    raw = run_bounded([sys.executable, "-c", command], ProcessLimits(15, 512 * MIB), input_bytes=payload)
    assert json.loads(raw) == {"result": "Converted document"}


def test_sparse_workbook_rejected_before_dense_conversion(tmp_path):
    path = workbook_path(tmp_path)
    assert path.stat().st_size < 10_000
    office_processor._validate_document_archive(path)
    with pytest.raises(ValueError, match="logical cell"):
        office_processor._validate_spreadsheet(path)


def test_renamed_spreadsheet_cannot_bypass_logical_limits(tmp_path):
    path = workbook_path(tmp_path).rename(tmp_path / "table.html")
    office_processor._validate_document_archive(path)
    with pytest.raises(ValueError, match="logical cell"):
        office_processor._validate_spreadsheet(path)


def test_forged_dimension_cannot_hide_sparse_cell(tmp_path):
    path = workbook_path(tmp_path)
    original = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(path, "w") as output:
        for item in source.infolist():
            data = source.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'ref="A1:DX1000"', b'ref="A1:A1"')
            output.writestr(item, data)
    with pytest.raises(ValueError, match="logical cell"):
        office_processor._validate_spreadsheet(path)


def test_sheet_relationship_resolved_before_budgeting(tmp_path):
    path = workbook_path(tmp_path)
    original = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(path, "w") as output:
        for item in source.infolist():
            data = source.read(item)
            name = item.filename
            if name == "xl/_rels/workbook.xml.rels":
                data = data.replace(b"/xl/worksheets/sheet1.xml", b"/xl/alternative.xml")
            if name == "xl/worksheets/sheet1.xml":
                name = "xl/alternative.xml"
            output.writestr(name, data)
    with pytest.raises(ValueError, match="logical cell"):
        office_processor._validate_spreadsheet(path)


@pytest.mark.parametrize("rows,columns", [(10_001, 1), (1, 257)])
def test_spreadsheet_coordinate_limits(tmp_path, rows, columns):
    with pytest.raises(ValueError, match="row or column"):
        office_processor._validate_spreadsheet(workbook_path(tmp_path, rows=rows, columns=columns))


def test_spreadsheet_sheet_and_combined_cell_budgets(tmp_path, monkeypatch):
    path = workbook_path(tmp_path, rows=20, columns=20, sheets=2)
    monkeypatch.setattr(office_processor, "MAX_SPREADSHEET_CELLS", 500)
    with pytest.raises(ValueError, match="logical cell"):
        office_processor._validate_spreadsheet(path)
    monkeypatch.setattr(office_processor, "MAX_SPREADSHEET_SHEETS", 1)
    with pytest.raises(ValueError, match="sheet count"):
        office_processor._validate_spreadsheet(path)


def test_legacy_xls_uses_same_logical_budget(monkeypatch, tmp_path):
    workbook = MagicMock()
    workbook.nsheets = 1
    workbook.sheet_by_index.return_value = SimpleNamespace(nrows=1000, ncols=128)
    manager = MagicMock()
    manager.__enter__.return_value = workbook
    monkeypatch.setitem(sys.modules, "xlrd", SimpleNamespace(open_workbook=lambda *args, **kwargs: manager))
    with pytest.raises(ValueError, match="logical cell"):
        office_processor._validate_spreadsheet(tmp_path / "table.xls")


def test_small_spreadsheet_still_converts_in_isolated_worker(tmp_path):
    chunks = office_processor.process_office(workbook_path(tmp_path, rows=3, columns=2), {})
    assert chunks and "Tail" in " ".join(chunk.text for chunk in chunks)


@pytest.mark.parametrize("stream", ["stdout", "stderr"])
def test_subprocess_pipe_output_is_bounded(stream):
    limits = ProcessLimits(seconds=5, memory_bytes=128 * MIB, stdout_bytes=1024, stderr_bytes=1024)
    with pytest.raises(ProcessingLimitError, match="output"):
        run_bounded([sys.executable, "-c", f"import sys; sys.{stream}.write('x'*16384)"], limits)


def test_subprocess_deadline_terminates_worker():
    started = time.monotonic()
    with pytest.raises(ProcessingLimitError, match="time limit"):
        run_bounded([sys.executable, "-c", "import time; time.sleep(10)"], ProcessLimits(0.1, 128 * MIB))
    assert time.monotonic() - started < 3


def test_subprocess_memory_watchdog_terminates_worker():
    with pytest.raises(ProcessingLimitError, match="memory limit"):
        run_bounded([sys.executable, "-c", "import time; data=bytearray(80*1024*1024); time.sleep(5)"],
                    ProcessLimits(10, 64 * MIB))


def test_parser_child_does_not_inherit_credentials(monkeypatch):
    monkeypatch.setenv("OPENROUTER_API_KEY", "test-only-secret")
    monkeypatch.setenv("LUMINRAG_API_TOKEN", "test-only-token")
    raw = run_bounded([sys.executable, "-c", "import json,os; print(json.dumps([key for key in os.environ if key in ('OPENROUTER_API_KEY','LUMINRAG_API_TOKEN')]))"],
                      ProcessLimits(5, 128 * MIB))
    assert json.loads(raw) == []


def test_parser_disables_onnx_telemetry_before_native_import():
    raw = run_bounded([sys.executable, "-c", "import os; print(os.environ.get('ORT_DISABLE_TELEMETRY'))"],
                      ProcessLimits(5, 128 * MIB))
    assert raw.strip() == b"1"


def test_ffmpeg_has_duration_size_protocol_and_process_limits(tmp_path, monkeypatch):
    source = small_wav(tmp_path / "source.wav")
    output_dir = tmp_path / "decoded"
    output_dir.mkdir()
    called = {}
    def fake_run(command, limits):
        called.update(command=command, limits=limits)
        small_wav(Path(command[-1]))
    monkeypatch.setattr(media_limits, "run_bounded", fake_run)
    result = media_limits.decode_audio(source, str(output_dir))
    assert result.exists()
    command = called["command"]
    assert command[command.index("-protocol_whitelist") + 1] == "file"
    assert command[command.index("-t") + 1] == str(media_limits.MAX_MEDIA_SECONDS + 1)
    assert command[command.index("-fs") + 1] == str(media_limits.MAX_WAV_BYTES)
    assert called["limits"].seconds == 120


def test_ffmpeg_failure_removes_partial_output(tmp_path, monkeypatch):
    source = small_wav(tmp_path / "source.wav")
    output_dir = tmp_path / "decoded"
    output_dir.mkdir()
    def fail(command, limits):
        Path(command[-1]).write_bytes(b"partial")
        raise ProcessingLimitError("time limit")
    monkeypatch.setattr(media_limits, "run_bounded", fail)
    with pytest.raises(ProcessingLimitError):
        media_limits.decode_audio(source, str(output_dir))
    assert not (output_dir / "audio.wav").exists()


def test_whisper_only_receives_its_model_settings(tmp_path, monkeypatch):
    call = MagicMock(return_value=[])
    monkeypatch.setattr(media_limits, "run_parser", call)
    media_limits.transcribe_bounded(tmp_path / "audio.wav", {
        "whisper": {"model_size": "tiny", "api_key": "test-only-secret"},
        "generation": {"api_key": "another-test-secret"},
    })
    assert call.call_args.kwargs["options"] == {"model_size": "tiny", "device": "cpu", "compute_type": "int8"}


def test_whisper_receives_normalized_pcm_without_redecoding_through_pyav(tmp_path, monkeypatch):
    source = tmp_path / "canonical.wav"
    pcm = np.array([-32768, -16384, 0, 16384, 32767], dtype="<i2")
    with wave.open(str(source), "wb") as output:
        output.setparams((1, 2, 16_000, 0, "NONE", "not compressed"))
        output.writeframes(pcm.tobytes())

    def transcribe(samples, **_settings):
        # File inputs invoke faster-whisper's incompatible PyAV decoder;
        # its supported waveform input must preserve PCM signs and scale.
        assert isinstance(samples, np.ndarray) and samples.dtype == np.float32
        np.testing.assert_allclose(samples, [-1.0, -0.5, 0.0, 0.5, 32767 / 32768])
        return iter([SimpleNamespace(start=0.0, end=0.001, text=" Synthetic speech. ")]), None

    model = SimpleNamespace(transcribe=transcribe)
    monkeypatch.setitem(sys.modules, "faster_whisper", SimpleNamespace(WhisperModel=lambda *args, **kwargs: model))
    assert media_limits.transcribe_local(source, {"model_size": "small"}) == [
        {"start": 0.0, "end": 0.001, "text": "Synthetic speech."},
    ]


@pytest.mark.parametrize("segments,error", [
    ([SimpleNamespace(start=0, end=float("inf"), text="speech")], "timestamps"),
    ([SimpleNamespace(start=0, end=1, text="12345")], "text"),
    ([SimpleNamespace(start=0, end=1, text="a")] * 3, "segment count"),
])
def test_whisper_iteration_bounds(tmp_path, monkeypatch, segments, error):
    model = MagicMock()
    model.transcribe.return_value = (iter(segments), None)
    monkeypatch.setitem(sys.modules, "faster_whisper", SimpleNamespace(WhisperModel=lambda *args, **kwargs: model))
    monkeypatch.setattr(media_limits, "MAX_TRANSCRIPT_CHARS", 4)
    monkeypatch.setattr(media_limits, "MAX_TRANSCRIPT_SEGMENTS", 2)
    with pytest.raises(ValueError, match=error):
        media_limits.transcribe_local(small_wav(tmp_path / "audio.wav"), {"model_size": "tiny"})


@pytest.mark.parametrize("window", [0, -1, float("inf"), float("nan")])
def test_audio_chunk_window_cannot_create_unbounded_loop(window):
    with pytest.raises(ValueError, match="window"):
        _build_chunks("source", [], window)


def fake_capture(monkeypatch, *, width=16, height=16, count=3, fps=1):
    cap = MagicMock()
    values = {video_transcriber.cv2.CAP_PROP_FPS: fps, video_transcriber.cv2.CAP_PROP_FRAME_COUNT: count,
              video_transcriber.cv2.CAP_PROP_FRAME_WIDTH: width, video_transcriber.cv2.CAP_PROP_FRAME_HEIGHT: height}
    cap.get.side_effect = lambda key: values[key]
    cap.read.side_effect = [(True, np.full((16, 16, 3), value, dtype=np.uint8)) for value in (0, 255, 0)] + [(False, None)]
    monkeypatch.setattr(video_transcriber.cv2, "VideoCapture", lambda *args: cap)
    return cap


def test_video_rejects_oversized_metadata_before_decode(tmp_path, monkeypatch):
    cap = fake_capture(monkeypatch, width=50_000, height=50_000)
    with pytest.raises(ValueError, match="dimensions"):
        video_transcriber._extract_keyframes_local(tmp_path / "video.mp4", tmp_path, .85)
    cap.read.assert_not_called()
    cap.release.assert_called_once()


def test_video_keyframes_and_output_bytes_are_bounded(tmp_path, monkeypatch):
    cap = fake_capture(monkeypatch)
    monkeypatch.setattr(video_transcriber, "MAX_KEYFRAMES", 1)
    with pytest.raises(ValueError, match="keyframe count"):
        video_transcriber._extract_keyframes_local(tmp_path / "video.mp4", tmp_path, .85)
    cap.release.assert_called_once()
    cap = fake_capture(monkeypatch)
    monkeypatch.setattr(video_transcriber, "MAX_KEYFRAME_BYTES", 1)
    with pytest.raises(ValueError, match="size limit"):
        video_transcriber._extract_keyframes_local(tmp_path / "video.mp4", tmp_path, .85)
    cap.release.assert_called_once()


def test_video_actual_frame_count_is_checked_when_metadata_unknown(tmp_path, monkeypatch):
    cap = fake_capture(monkeypatch, count=0)
    monkeypatch.setattr(video_transcriber, "MAX_VIDEO_FRAMES", 2)
    with pytest.raises(ValueError, match="frames or duration"):
        video_transcriber._extract_keyframes_local(tmp_path / "video.mp4", tmp_path, .85)
    cap.release.assert_called_once()
