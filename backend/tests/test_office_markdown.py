from pathlib import Path
import time
import zipfile

import pytest

from backend.ingestion.markdown_processor import chunk_markdown
from backend.ingestion.markdown_processor import _heading_title
from backend.ingestion.office_processor import _validate_document_archive
from backend.ingestion.office_processor import process_office
from backend.utils.uploads import UploadRejected, validate_file


def test_markdown_headings_and_overflow_share_pdf_chunking():
    chunks = chunk_markdown("# First\n" + "Words. " * 100 + "\n# Second\nAnother idea.", "course", {"markdown_processor": {"max_chunk_chars": 250}})
    assert len(chunks) > 2
    assert all(len(chunk.text) <= 250 for chunk in chunks)
    assert chunks[-1].metadata["section_title"] == "Second"
    assert all(chunk.modality == "markdown" for chunk in chunks)


def test_docx_and_pptx_use_off_the_shelf_conversion(tmp_path):
    from docx import Document
    from pptx import Presentation
    document = Document()
    document.add_heading("Biology", level=1)
    document.add_paragraph("Photosynthesis uses light energy.")
    word = tmp_path / "lesson.docx"
    document.save(word)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Biology"
    slide.placeholders[1].text = "Photosynthesis uses light energy."
    slides = tmp_path / "lesson.pptx"
    presentation.save(slides)
    for path in (word, slides):
        chunks = process_office(path, {})
        assert chunks and "Photosynthesis" in " ".join(chunk.text for chunk in chunks)
        assert all(chunk.modality == "office" for chunk in chunks)


def test_hosted_rejects_audio_and_accepts_office_mime(monkeypatch):
    monkeypatch.setenv("LUMIN_HOSTED", "1")
    with pytest.raises(UploadRejected) as error:
        validate_file("lecture.mp3", "audio/mpeg")
    assert error.value.status == 415
    assert validate_file("lesson.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document") == "lesson.docx"
    assert validate_file("lesson.html", "text/html") == "lesson.html"


def test_pathological_markdown_heading_is_linear():
    malicious = "# " + " " * 2_000_000 + "#"
    started = time.perf_counter()
    assert _heading_title(malicious) is None
    assert time.perf_counter() - started < 1.0


def test_office_xml_parser_rejects_utf16_entities_but_accepts_epub_doctype(tmp_path):
    malicious = tmp_path / "unsafe.docx"
    xml = '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE doc [<!ENTITY leak SYSTEM "file:///secret">]><doc>&leak;</doc>'
    with zipfile.ZipFile(malicious, "w") as archive:
        archive.writestr("word/document.xml", xml.encode("utf-16"))
    with pytest.raises(ValueError, match="Unsafe"):
        _validate_document_archive(malicious)
    safe = tmp_path / "safe.epub"
    with zipfile.ZipFile(safe, "w") as archive:
        archive.writestr("chapter.xhtml", '<!DOCTYPE html><html xmlns="http://www.w3.org/1999/xhtml"><body>A chapter</body></html>')
    _validate_document_archive(safe)
